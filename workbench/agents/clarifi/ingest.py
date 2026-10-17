"""Bounded multi-table CSV and JSON intake shared by CLI and UI."""
from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path

from .schemas import InvestigationInput

TABLES = ("employees", "accounts", "relationships", "trades", "events", "entities", "payments")
TOOL_TABLES = {
    "insider_trading": ("employees", "accounts", "relationships", "trades", "events"),
    "shell_companies": ("entities", "payments"),
    "combined": TABLES,
}


class IntakeError(ValueError):
    """Reviewed input guidance safe to display without echoing source records."""


def excel_template(tool: str) -> bytes:
    """An editable example with one sheet per evidence table; never real records."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    data = sample_input(tool)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for table in TOOL_TABLES[tool]:
        sheet = workbook.create_sheet(table)
        rows = [row.model_dump(mode="json") for row in getattr(data, table)]
        record_type = InvestigationInput.model_fields[table].annotation.__args__[0]
        headers = list(record_type.model_fields)
        sheet.append(headers)
        for row in rows:
            values = [row.get(key) for key in headers]
            values = ["|".join(value) if isinstance(value, list) else value for value in values]
            sheet.append(values)
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="0057B8")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(48, max(20, len(str(column[0].value)) + 4))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def read_excel_bytes(content: bytes, tool: str) -> dict:
    """Read bounded .xlsx tables through the same normalization as CSV intake."""
    from openpyxl import load_workbook
    if tool not in TOOL_TABLES or len(content) > 5 * 1024 * 1024:
        raise IntakeError("Invalid scope or workbook exceeds 5 MiB")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
            raise IntakeError("Expanded workbook exceeds 30 MiB")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        names = set(workbook.sheetnames)
        if names - set(TOOL_TABLES[tool]):
            if tool != "combined" and names <= set(TABLES) and {"employees", "accounts", "trades", "entities", "payments"} <= names:
                raise IntakeError("This Excel workbook contains both trading and company-payment records. Select Combined investigation under ClariFi Capabilities, then choose Upload Excel workbook and upload the same file. No spreadsheet edits are needed.")
            raise IntakeError("Workbook contains sheets outside the selected capability. Use Combined investigation for a workbook with both trading and company-payment tables, or use the selected capability's Excel template.")
        tables = {}
        for sheet in workbook:
            if (sheet.max_row or 0) > 5001 or (sheet.max_column or 0) > 30:
                raise IntakeError("Workbook table exceeds supported dimensions")
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            headers = []
            for index, cells in enumerate(sheet.iter_rows()):
                if index > 5000 or len(cells) > 30:
                    raise IntakeError("Workbook table exceeds supported dimensions")
                if any(cell.data_type == "f" for cell in cells):
                    raise IntakeError("Use values rather than formulas")
                values = [cell.value for cell in cells]
                if index == 0:
                    headers = values
                else:
                    # Some Excel writers persist the text-entry apostrophe.
                    # Strip it only for timestamp columns; timezone validation
                    # still applies and native timezone-less dates stay invalid.
                    values = [value[1:] if column < len(headers) and headers[column] in {"occurred_at", "public_at"}
                              and isinstance(value, str) and value.startswith("'") else value
                              for column, value in enumerate(values)]
                if any(value is not None for value in values):
                    writer.writerow(values)
            if output.tell():
                tables[sheet.title] = read_csv_bytes(output.getvalue().encode())
        return tables
    finally:
        workbook.close()


def read_csv_bytes(content: bytes):
    if len(content) > 5 * 1024 * 1024:
        raise IntakeError("CSV exceeds 5 MiB")
    reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig"), newline=""))
    if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
        raise IntakeError("CSV requires unique column headers")
    rows = []
    for raw in reader:
        if None in raw or any(value is None for value in raw.values()):
            raise IntakeError("CSV row does not match its headers")
        row = {key: value.strip() for key, value in raw.items() if value.strip()}
        if "access_issuers" in row:
            row["access_issuers"] = [value.strip() for value in row["access_issuers"].split("|") if value.strip()]
        if "precleared" in row:
            if row["precleared"].lower() not in {"true", "false"}:
                raise IntakeError("precleared must be true, false, or empty")
            row["precleared"] = row["precleared"].lower() == "true"
        rows.append(row)
        if len(rows) > 5000:
            raise IntakeError("CSV exceeds 5,000 rows; narrow the scope")
    return rows


def sample_input(tool: str) -> InvestigationInput:
    if tool not in TOOL_TABLES:
        raise IntakeError("Unknown financial capability")
    sample = Path(__file__).resolve().parents[3] / "examples/clarifi/financial_case.json"
    data = json.loads(sample.read_text(encoding="utf-8"))
    return InvestigationInput.model_validate({"case_id": f"synthetic-{tool}",
                                             **{key: data[key] for key in TOOL_TABLES[tool]}})


def validate_scope(data: InvestigationInput, tool: str) -> InvestigationInput:
    expected = ["insider_trading", "shell_companies"] if tool == "combined" else [tool]
    if data.tools != expected:
        raise IntakeError("Supply the tables for the selected capability. Choose Combined investigation for both sets of tables.")
    return data
