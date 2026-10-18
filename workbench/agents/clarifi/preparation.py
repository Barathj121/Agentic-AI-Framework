"""Local, approval-gated migration → sandboxed ClariFi handoff.

Uses MigrationService for each related table; never executes generated code here.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path

from filelock import FileLock

from workbench.data import decode_csv
from workbench.migration import MetadataField, MigrationService, _metadata_csv
from workbench.query import readonly_connection
from workbench.workspace import WorkspaceManager, now, read_json, safe_path
from .ingest import IntakeError, TABLES, TOOL_TABLES, read_csv_bytes, validate_scope
from .schemas import InvestigationInput

MAX_BYTES = 5 * 1024 * 1024
REQUIRED = {"insider_trading": {"employees", "accounts", "trades"},
            "shell_companies": {"entities", "payments"},
            "combined": {"employees", "accounts", "trades", "entities", "payments"}}


def requests_clarifi(instruction: str) -> bool:
    """Explicit routing only; mentioning a financial topic doesn't select an agent."""
    return bool(re.search(r"\buse\s+(?:the\s+)?(?:clari\s*fi|clarify)\s+agent\b", instruction, re.I))


def digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


@dataclass(frozen=True)
class SourceTable:
    key: str
    filename: str
    label: str
    content: bytes
    original_sha256: str


def records_csv(rows: list[dict]) -> bytes:
    if not rows or len(rows) > 5000 or any(not isinstance(row, dict) for row in rows):
        raise IntakeError("Each source table needs 1–5,000 object records")
    columns = list(dict.fromkeys(key for row in rows for key in row))
    if not columns or len(columns) > 30 or any(not isinstance(key, str) or not key.strip() for key in columns):
        raise IntakeError("Source tables need 1–30 named columns")
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=columns)
    writer.writeheader()
    for row in rows:
        converted = {}
        for key, value in row.items():
            if isinstance(value, list) and all(isinstance(item, str) and "|" not in item for item in value):
                value = "|".join(value)
            elif isinstance(value, (dict, list)):
                raise IntakeError("Flatten nested objects into separate tables before uploading")
            if isinstance(value, bool):
                value = str(value).lower()
            converted[key] = value
        writer.writerow(converted)
    return output.getvalue().encode("utf-8")


def checked_csv(content: bytes) -> bytes:
    text, _, _ = decode_csv(content)
    try:
        delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not rows or not 1 <= len(rows[0]) <= 30 or not 2 <= len(rows) <= 5001:
        raise IntakeError("Source tables need headers, 1–5,000 rows and at most 30 columns")
    headers = rows[0]
    if any(not key.strip() or len(key) > 200 for key in headers) or len({key.casefold() for key in headers}) != len(headers):
        raise IntakeError("Use unique, nonempty source headers")
    if any(len(row) != len(headers) for row in rows[1:]):
        raise IntakeError("Source rows must match the header columns")
    output = io.StringIO(newline="")
    csv.writer(output).writerows(rows)
    return output.getvalue().encode("utf-8")


def read_sources(files: list[tuple[str, bytes]]) -> list[SourceTable]:
    """Structured input only; original file hashes survive normalization."""
    if not files or len(files) > 7 or sum(len(content) for _, content in files) > MAX_BYTES:
        raise IntakeError("Upload 1–7 CSV, Excel or JSON files, at most 5 MiB combined")
    result = []
    for filename, content in files:
        if not filename or Path(filename).name != filename or "\\" in filename or ":" in filename:
            raise IntakeError("Use a filename without a directory path")
        suffix = Path(filename).suffix.lower()
        tables = {}
        if suffix == ".csv":
            tables[Path(filename).stem] = checked_csv(content)
        elif suffix == ".json":
            try:
                document = json.loads(content)
            except (ValueError, UnicodeError):
                raise IntakeError("JSON must be a table of objects or an object containing tables") from None
            if isinstance(document, list):
                document = {Path(filename).stem: document}
            if not isinstance(document, dict) or set(document) - {"case_id"} == set():
                raise IntakeError("JSON must contain source tables")
            for label, rows in document.items():
                if label != "case_id" and rows != []:
                    tables[label] = records_csv(rows) if isinstance(rows, list) else records_csv([])
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 200 or sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise IntakeError("Expanded workbook exceeds supported limits")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.sheetnames) > 7:
                    raise IntakeError("Upload at most seven source worksheets")
                for sheet in book:
                    if (sheet.max_row or 0) > 5001 or (sheet.max_column or 0) > 30:
                        raise IntakeError("Worksheet exceeds supported dimensions")
                    output = io.StringIO(newline="")
                    writer = csv.writer(output)
                    headers = []
                    for index, cells in enumerate(sheet.iter_rows()):
                        if index > 5000 or len(cells) > 30:
                            raise IntakeError("Worksheet exceeds supported dimensions")
                        if any(cell.data_type == "f" for cell in cells):
                            raise IntakeError("Use values rather than Excel formulas")
                        values = [cell.value.isoformat() if hasattr(cell.value, "isoformat") else cell.value for cell in cells]
                        if index == 0:
                            headers = values
                        else:
                            # Match direct ClariFi intake: some workbook writers
                            # retain Excel's timestamp text-entry apostrophe.
                            values = [value[1:] if column < len(headers) and headers[column] in {"occurred_at", "public_at"}
                                      and isinstance(value, str) and value.startswith("'") else value
                                      for column, value in enumerate(values)]
                        if any(value is not None for value in values):
                            writer.writerow(values)
                    if len(output.getvalue().splitlines()) > 1:
                        tables[sheet.title] = checked_csv(output.getvalue().encode("utf-8"))
            finally:
                book.close()
        else:
            raise IntakeError("Use CSV, Excel or JSON. PDF extraction is not supported by this local flow yet.")
        for label, table in tables.items():
            if len(result) >= 7:
                raise IntakeError("Upload at most seven related source tables")
            result.append(SourceTable(str(len(result)), filename, str(label), table, digest(content)))
    if not result or sum(len(item.content) for item in result) > MAX_BYTES:
        raise IntakeError("Normalized source tables exceed 5 MiB or contain no data")
    return result


def target_metadata(table: str) -> bytes:
    record = InvestigationInput.model_fields[table].annotation.__args__[0]
    fields = []
    for name, field in record.model_fields.items():
        meaning = {
            "entity_id": "Unique identifier of this company itself, not its owner",
            "owner_entity_id": "Identifier of another company that owns this company; only with documented ownership",
            "employee_id": "Employee identifier; a unique staff ID in employees, an employee reference in other tables",
            "account_id": "Trading account identifier; a unique account ID in accounts, an account reference in other tables",
            "sender_id": "Entity identifier of the company paying the money",
            "receiver_id": "Entity identifier of the company receiving the money",
            "name": "Full employee name" if table == "employees" else "Registered company name",
            "owner_name": "Name of the trading account holder, not necessarily an employee",
            "transaction_id": "Unique payment record identifier",
            "trade_id": "Unique trading record identifier, not the trading account identifier",
        }.get(name, name.replace("_", " "))
        hint = "Preserve source meaning; never invent identities, relationships, values or missing evidence."
        if name in {"occurred_at", "public_at"}:
            hint += " ISO 8601 text with explicit UTC offset; do not infer a timezone."
        if name == "access_issuers":
            hint += " Pipe-separated issuer IDs."
        if name == "currency":
            hint += " Three-letter uppercase currency; do not convert monetary values."
        fields.append(MetadataField(column_name=name, data_type="VARCHAR", required=field.is_required(),
                                    description=f"ClariFi {table}.{name}: {meaning}. {hint}", transformation_hint=hint))
    return _metadata_csv(fields).encode()


class FinancialPreparation:
    def __init__(self, settings, *, llm=None, sandbox=None):
        self.settings = settings
        self.manager = WorkspaceManager(settings.workspace_root / "clarifi-pipelines")
        self.llm, self.sandbox = llm, sandbox

    def workspace(self, identifier):
        workspace = self.manager.get(identifier)
        if self.manager.read(workspace, "metadata/workspace.json").get("application") != "clarifi_preparation":
            raise ValueError("Not a ClariFi preparation workspace")
        return workspace

    def migration(self, identifier):
        return MigrationService(replace(self.settings, workspace_root=self.workspace(identifier) / "migrations"),
                                llm=self.llm, sandbox=self.sandbox)

    def state(self, identifier):
        return self.manager.read(self.workspace(identifier), "metadata/preparation.json")

    def propose(self, identifier, table):
        item = self.state(identifier)["tables"][table]
        plan = self.migration(identifier).propose(item["workspace_id"], exact_targets=True)
        if self.llm is not None:
            self.manager.write(self.workspace(identifier), "metadata/mapping_model_" + table + ".json",
                               {"provider": type(self.llm).__name__, "model": getattr(self.llm, "model", None)})
        return plan

    def create(self, case_id, instruction, tool, sources, assignments):
        if not requests_clarifi(instruction):
            raise IntakeError('Include "use ClariFi agent" in your instruction to select this flow')
        # Validate the case identifier before any model or sandbox call.
        from pydantic import TypeAdapter
        from .schemas import Identifier
        TypeAdapter(Identifier).validate_python(case_id)
        if tool not in REQUIRED or set(assignments) != {item.key for item in sources}:
            raise IntakeError("Assign every source table to its ClariFi evidence table")
        selected = list(assignments.values())
        if len(selected) != len(set(selected)) or not set(selected) <= set(TOOL_TABLES[tool]) or not REQUIRED[tool] <= set(selected):
            raise IntakeError("Provide each required table once. Combine split files for the same table before uploading.")
        workspace = self.manager.create(case_id, instruction)
        self.manager.status(workspace, "PREPARING", application="clarifi_preparation", kind="clarifi")
        migration = self.migration(workspace.name)
        state = {"case_id": case_id, "instruction": instruction, "tool": tool, "tables": {}}
        try:
            for item in sources:
                table = assignments[item.key]
                headers = next(csv.reader(io.StringIO(item.content.decode("utf-8"))))
                metadata = _metadata_csv([MetadataField(column_name=key, data_type="VARCHAR", required=False,
                    description="Original source text. Business meaning requires review.") for key in headers]).encode()
                child = migration.create(f"{case_id} · {table}", item.content, target_metadata(table), metadata,
                    problem=f"{instruction}\nMap only this source table to ClariFi {table}. Preserve record count, identifiers and timezone offsets. Do not invent business evidence.",
                    preserve_source_text=True)
                state["tables"][table] = {"workspace_id": child.name, "filename": item.filename, "source_table": item.label,
                    "original_sha256": item.original_sha256, "normalized_sha256": digest(item.content)}
            self.manager.write(workspace, "metadata/preparation.json", state)
            self.manager.status(workspace, "MAPPING_REVIEW")
            return workspace
        except Exception:
            self.manager.status(workspace, "PREPARATION_FAILED")
            raise

    def prepare(self, identifier):
        workspace = self.workspace(identifier)
        with FileLock(str(safe_path(workspace, "metadata/workflow.lock")), timeout=0):
            state, migration = self.state(identifier), self.migration(identifier)
            for item in state["tables"].values():
                child = migration.workspace(item["workspace_id"])
                if not read_json(child / "metadata/mapping_approval.json", {}).get("human_approved"):
                    raise IntakeError("Review and approve mappings for every source table before conversion")
            for name in ("prepared.json", "handoff.json", "investigation.json"):
                safe_path(workspace, "results/" + name).unlink(missing_ok=True)
            self.manager.status(workspace, "MIGRATING")
            try:
                payload, evidence = {"case_id": state["case_id"]}, {}
                for table, item in state["tables"].items():
                    child = migration.workspace(item["workspace_id"])
                    report = migration.execute(child.name)
                    if not report.get("complete"):
                        raise IntakeError("Migration is incomplete. Correct missing fields or invalid records in a new preparation before investigating.")
                    content = safe_path(child, "results/migrated_data.csv").read_bytes()
                    payload[table] = read_csv_bytes(content)
                    # Independently bind the handoff to the materialized table,
                    # rather than accepting a generated/exported CSV alone.
                    with readonly_connection(child) as connection:
                        cursor = connection.execute('SELECT * FROM results.migrated_data ORDER BY rowid')
                        columns = [column[0] for column in cursor.description]
                        materialized = [dict(zip(columns, row)) for row in cursor.fetchall()]
                    if payload[table] != read_csv_bytes(records_csv(materialized)):
                        raise IntakeError("Migration export disagrees with its validated table; investigation is blocked")
                    evidence[table] = {**item, "target_sha256": digest(content),
                        "mapping_sha256": digest(safe_path(child, "metadata/mapping_approval.json").read_bytes()),
                        "rows": len(payload[table])}
                source = validate_scope(InvestigationInput.model_validate(payload), state["tool"])
                content = source.model_dump_json(indent=2).encode()
                if len(content) > MAX_BYTES:
                    raise IntakeError("Prepared investigation exceeds 5 MiB")
                safe_path(workspace, "results/prepared.json").write_bytes(content)
                self.manager.write(workspace, "results/handoff.json", {"prepared_sha256": digest(content), "tables": evidence,
                                   "prepared_at": now(), "execution": "docker_sandbox"})
                self.manager.status(workspace, "READY_FOR_INVESTIGATION")
                return source
            except Exception:
                self.manager.status(workspace, "PREPARATION_FAILED")
                raise

    def prepared(self, identifier):
        workspace = self.workspace(identifier)
        handoff = read_json(safe_path(workspace, "results/handoff.json"))
        if not handoff:
            raise IntakeError("Complete migration and validation before investigation")
        content = safe_path(workspace, "results/prepared.json").read_bytes()
        if digest(content) != handoff["prepared_sha256"]:
            raise IntakeError("Prepared data changed; prepare and review it again")
        migration = self.migration(identifier)
        for item in handoff["tables"].values():
            child = migration.workspace(item["workspace_id"])
            if digest(safe_path(child, "results/migrated_data.csv").read_bytes()) != item["target_sha256"] or digest(safe_path(child, "metadata/mapping_approval.json").read_bytes()) != item["mapping_sha256"]:
                raise IntakeError("Migration evidence changed; prepare and review it again")
        return validate_scope(InvestigationInput.model_validate_json(content), self.state(identifier)["tool"])

    def investigate(self, identifier, *, approved=False, mode="local", model=None):
        if approved is not True:
            raise IntakeError("Review the converted data and approve this investigation")
        if mode not in {"local", "gemini", "vertex", "rules"}:
            raise IntakeError("This local pipeline requires a local ClariFi orchestrator with Docker access")
        workspace = self.workspace(identifier)
        with FileLock(str(safe_path(workspace, "metadata/workflow.lock")), timeout=0):
            source = self.prepared(identifier)
            safe_path(workspace, "results/investigation.json").unlink(missing_ok=True)
            self.manager.write(workspace, "metadata/investigation_approval.json", {"approved_at": now(), "mode": mode,
                "model": model, "prepared_sha256": digest(safe_path(workspace, "results/prepared.json").read_bytes())})
            self.manager.status(workspace, "INVESTIGATING")
            try:
                from .agent import ClariFiAgent
                from .sandbox import FinancialSandbox
                executor = FinancialSandbox(self.settings, workspace, sandbox=self.sandbox)
                report = ClariFiAgent(mode=mode, model=model, tool_executor=executor).run(source, approved=True)
                handoff = read_json(safe_path(workspace, "results/handoff.json"))
                report["provenance"]["migration"] = handoff
                report["trace"].insert(0, {"stage": "migration_handoff", "tables": list(handoff["tables"]), "passed": True})
                self.manager.write(workspace, "results/investigation.json", {"input": source.model_dump(mode="json"), "report": report})
                self.manager.status(workspace, "COMPLETED")
                return report
            except Exception:
                self.manager.status(workspace, "INVESTIGATION_FAILED")
                raise


def canonical_input(sources, assignments, *, case_id, tool):
    """Accept only an exact existing contract, without generated mappings or guessed values."""
    if set(assignments) != {item.key for item in sources}:
        raise IntakeError("Assign every source table before checking the ClariFi format")
    selected = list(assignments.values())
    if len(selected) != len(set(selected)) or not set(selected) <= set(TOOL_TABLES[tool]):
        raise IntakeError("Each evidence table must be supplied once within the selected capability")
    tables = {assignments[item.key]: read_csv_bytes(item.content) for item in sources}
    return validate_scope(InvestigationInput.model_validate({"case_id": case_id, **tables}), tool)
